"""
export_helper.py — Helper generate file PDF dan Excel untuk laporan belanja.
Dipanggil oleh export_controller.py. Tidak ada logika bisnis di sini — hanya
formatting dan penulisan file.

Spesifikasi: PRD §11.3, Design System §10.2–§10.5
"""

from __future__ import annotations

from datetime import datetime
from typing import Any

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph,
    Spacer, HRFlowable,
)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER  # type: ignore[attr-defined]
from reportlab.pdfgen.canvas import Canvas

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from utils.format_helper import format_rupiah, format_jumlah


# ─── TOKEN WARNA (Design System §10.2 & §10.3) ───────────────────────────────

def _hex_ke_rgb(hex_str: str) -> colors.Color:  # type: ignore[name-defined]
    """Konversi hex color string ke Color object untuk reportlab."""
    hex_str = hex_str.lstrip("#")
    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
    return colors.Color(r / 255, g / 255, b / 255)  # type: ignore[attr-defined]


# Warna PDF (palet light mode — Design System §10.2)
PDF_TEAL_900 = _hex_ke_rgb("#134E4A")   # BG_SIDEBAR — header blok & ringkasan
PDF_TEAL_600 = _hex_ke_rgb("#0D9488")   # ACCENT_PRIMARY — judul & header kolom
PDF_TEAL_200 = _hex_ke_rgb("#99F6E4")   # BORDER_COLOR — garis tabel & divider
PDF_TEAL_100 = _hex_ke_rgb("#CCFBF1")   # BTN_SECONDARY_BG — footer total
PDF_TEAL_50  = _hex_ke_rgb("#F0FDFA")   # BG_TABLE_ROW_ALT — baris alternating
PDF_PUTIH    = _hex_ke_rgb("#FFFFFF")
PDF_TEKS     = _hex_ke_rgb("#0F2421")   # TEXT_PRIMARY
PDF_TEKS_ABU = _hex_ke_rgb("#6B7280")   # abu netral
PDF_WARNING  = _hex_ke_rgb("#D97706")   # COLOR_WARNING
PDF_DANGER   = _hex_ke_rgb("#DC2626")   # COLOR_DANGER

# Warna Excel — hex string tanpa "#" (format yang diterima openpyxl)
XL_TEAL_900 = "134E4A"
XL_TEAL_600 = "0D9488"
XL_TEAL_200 = "99F6E4"
XL_TEAL_100 = "CCFBF1"
XL_TEAL_50  = "F0FDFA"
XL_PUTIH    = "FFFFFF"
XL_DANGER   = "DC2626"
XL_WARNING  = "D97706"
XL_ABU      = "888888"
XL_TEKS     = "0F2421"

NAMA_BULAN = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]


# ─── HELPER BERSAMA ──────────────────────────────────────────────────────────

def _nama_file_default(data_periode: dict[str, Any]) -> str:
    """
    Buat nama file default sesuai pola Design System §10.5.
    Contoh: 'catatan_belanja_minggu4_november_2025'
    """
    nama_bulan = NAMA_BULAN[data_periode["bulan"] - 1].lower()
    return (
        f"catatan_belanja_minggu{data_periode['minggu_ke']}"
        f"_{nama_bulan}_{data_periode['tahun']}"
    )


def _warna_sisa_pdf(sisa: float, nominal: float) -> colors.Color:  # type: ignore[name-defined]
    """Pilih warna teks sisa budget sesuai threshold 20% (FR-DB-3)."""
    if sisa < 0:
        return PDF_DANGER
    if nominal > 0 and sisa < nominal * 0.20:
        return PDF_WARNING
    # Sisa sehat — teks putih agar kontras di atas bg teal-900
    return PDF_PUTIH


# ─── NUMBERED CANVAS (nomor halaman otomatis) ────────────────────────────────

class _NumberedCanvas(Canvas):
    """
    Canvas kustom yang menambahkan nomor halaman di footer kanan bawah.
    Hanya aktif jika laporan lebih dari 1 halaman.
    """

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        # Simpan snapshot state tiap halaman sebelum di-render
        self._saved_page_states: list[dict[str, Any]] = []

    def showPage(self) -> None:
        # Simpan seluruh state halaman saat ini ke list
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()  # type: ignore[attr-defined]

    def save(self) -> None:
        total_halaman = len(self._saved_page_states)
        for nomor, state in enumerate(self._saved_page_states, start=1):
            self.__dict__.update(state)
            if total_halaman > 1:
                # Cetak nomor halaman hanya jika multi-halaman
                self._cetak_nomor_halaman(nomor, total_halaman)
            super().showPage()
        super().save()

    def _cetak_nomor_halaman(self, nomor: int, total: int) -> None:
        """Tulis teks 'Halaman N dari T' di pojok kanan bawah."""
        self.saveState()
        self.setFont("Helvetica", 7.5)
        self.setFillColor(PDF_TEKS_ABU)
        teks = f"Halaman {nomor} dari {total}"
        self.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, teks)
        self.restoreState()


# ─── EKSPOR PDF ───────────────────────────────────────────────────────────────

def ekspor_pdf(path_file: str, data_periode: dict[str, Any], daftar_item: list[dict[str, Any]]) -> None:
    """
    Generate file PDF laporan belanja versi polished dan simpan ke path_file.

    data_periode berisi kunci:
        minggu_ke, bulan (1-12), tahun, tgl_mulai (yyyy-mm-dd),
        tgl_selesai (yyyy-mm-dd), nominal_budget, total_pengeluaran, sisa_budget

    daftar_item: list of dict dengan kunci:
        nama_barang, nama_kategori, jumlah_barang, nama_satuan,
        harga_satuan, harga_total
    """

    def _fmt(tanggal: str) -> str:
        """Konversi yyyy-mm-dd ke dd-mm-yyyy untuk tampilan."""
        return "-".join(reversed(tanggal.split("-")))

    nama_bulan   = NAMA_BULAN[data_periode["bulan"] - 1]
    tgl_mulai    = _fmt(data_periode["tgl_mulai"])
    tgl_selesai  = _fmt(data_periode["tgl_selesai"])
    sisa         = float(data_periode["sisa_budget"])
    nominal      = float(data_periode["nominal_budget"])
    total_keluar = float(data_periode["total_pengeluaran"])
    timestamp    = datetime.now().strftime("%d-%m-%Y %H:%M")

    # Margin dokumen A4
    M_L = M_R = 2.2 * cm
    M_T = 2.0 * cm
    M_B = 2.2 * cm

    doc = SimpleDocTemplate(
        path_file,
        pagesize=A4,
        leftMargin=M_L, rightMargin=M_R,
        topMargin=M_T, bottomMargin=M_B,
    )

    # Lebar area konten = lebar A4 dikurangi kedua margin
    lebar = A4[0] - M_L - M_R

    elemen: list[Any] = []

    # ── HELPER STYLE PARAGRAPH ────────────────────────────────────────────────

    def _gaya(nama: str, font: str = "Helvetica", ukuran: float = 9,
              warna: Any = PDF_TEKS, align: Any = TA_LEFT,
              leading: float = 0, space_before: float = 0, space_after: float = 0) -> ParagraphStyle:
        """Buat ParagraphStyle dengan parameter minimal."""
        return ParagraphStyle(
            nama,
            fontName=font,
            fontSize=ukuran,
            textColor=warna,
            alignment=align,
            leading=leading or ukuran * 1.3,
            spaceBefore=space_before,
            spaceAfter=space_after,
        )

    # ── 1. HEADER BLOK ────────────────────────────────────────────────────────
    # Satu card gelap teal-900 berisi judul + info periode

    gaya_judul = _gaya("judul", "Helvetica-Bold", 18, PDF_PUTIH, TA_LEFT, leading=22)
    gaya_sub   = _gaya("sub", "Helvetica", 9, PDF_TEAL_200, TA_LEFT, leading=13, space_before=3)
    gaya_tgl   = _gaya("tgl", "Helvetica", 8, PDF_TEAL_200, TA_LEFT, leading=12)

    isi_header: list[Any] = [
        Paragraph("Catatan Belanja", gaya_judul),
        Paragraph(
            f"Laporan Minggu ke-{data_periode['minggu_ke']} \u00b7 {nama_bulan} {data_periode['tahun']}",
            gaya_sub,
        ),
        Paragraph(f"{tgl_mulai}  \u2013  {tgl_selesai}", gaya_tgl),
    ]

    tbl_header = Table([[isi_header]], colWidths=[lebar])
    tbl_header.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), PDF_TEAL_900),
        ("TOPPADDING",    (0, 0), (-1, -1), 14),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
        ("LEFTPADDING",   (0, 0), (-1, -1), 18),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 18),
    ]))
    elemen.append(tbl_header)
    elemen.append(Spacer(1, 10))

    # ── 2. KARTU RINGKASAN STAT (3 kolom) ────────────────────────────────────
    # Tiga kartu teal-900 berdampingan dengan label kecil + nilai besar

    warna_sisa = _warna_sisa_pdf(sisa, nominal)
    GAP        = 5.0                          # jarak antar kartu dalam pt
    lebar_kol  = (lebar - 2 * GAP) / 3

    def _stat_isi(label: str, nilai: str, warna_nilai: Any = PDF_PUTIH) -> list[Any]:
        """Buat list [label, nilai] untuk satu kartu stat."""
        return [
            Paragraph(label, _gaya(f"sl_{label}", "Helvetica", 8,
                                   PDF_TEAL_200, TA_CENTER, space_after=4)),
            Paragraph(f"<b>{nilai}</b>", _gaya(f"sv_{label}", "Helvetica-Bold", 13,
                                               warna_nilai, TA_CENTER)),
        ]

    # Baris tabel stat: [kartu1, gap, kartu2, gap, kartu3]
    baris_stat = [[
        _stat_isi("Budget", format_rupiah(nominal)),
        "",
        _stat_isi("Total Pengeluaran", format_rupiah(total_keluar)),
        "",
        _stat_isi("Sisa Budget", format_rupiah(sisa), warna_sisa),
    ]]

    tbl_stat = Table(
        baris_stat,
        colWidths=[lebar_kol, GAP, lebar_kol, GAP, lebar_kol],
    )
    tbl_stat.setStyle(TableStyle([
        # Kartu di kolom 0, 2, 4 → background teal-900
        ("BACKGROUND",    (0, 0), (0, 0), PDF_TEAL_900),
        ("BACKGROUND",    (2, 0), (2, 0), PDF_TEAL_900),
        ("BACKGROUND",    (4, 0), (4, 0), PDF_TEAL_900),
        # Kolom gap (1, 3) → latar putih/transparan
        ("BACKGROUND",    (1, 0), (1, 0), PDF_PUTIH),
        ("BACKGROUND",    (3, 0), (3, 0), PDF_PUTIH),
        ("TOPPADDING",    (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING",   (0, 0), (-1, -1), 0),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 0),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elemen.append(tbl_stat)
    elemen.append(Spacer(1, 14))

    # ── 3. LABEL SEKSI TABEL ─────────────────────────────────────────────────
    elemen.append(Paragraph(
        "Daftar Belanja",
        _gaya("seksi", "Helvetica-Bold", 10, PDF_TEAL_600, TA_LEFT, space_after=6),
    ))

    # ── 4. TABEL ITEM BELANJA ─────────────────────────────────────────────────
    # Lebar kolom: No | Nama | Kategori | Jml & Sat | Harga/Sat | Total
    C_NO   = 0.8 * cm
    C_KAT  = 3.2 * cm
    C_JML  = 2.4 * cm
    C_HSAT = 2.8 * cm
    C_TOT  = 2.8 * cm
    C_NAMA = lebar - C_NO - C_KAT - C_JML - C_HSAT - C_TOT
    lebar_kolom = [C_NO, C_NAMA, C_KAT, C_JML, C_HSAT, C_TOT]

    def _ph(teks: str, align: int = TA_CENTER) -> Paragraph:  # type: ignore[assignment]
        """Paragraf header kolom tabel — putih di atas teal-600."""
        return Paragraph(teks, _gaya(f"ph_{teks}", "Helvetica-Bold", 8.5,
                                     PDF_PUTIH, align))

    def _pd(teks: str, align: int = TA_LEFT, bold: bool = False,  # type: ignore[assignment]
            warna: Any = PDF_TEKS) -> Paragraph:
        """Paragraf sel data tabel."""
        font = "Helvetica-Bold" if bold else "Helvetica"
        return Paragraph(str(teks), _gaya(f"pd_{teks[:8]}", font, 8.5, warna, align))

    # Baris header kolom
    baris_header = [[
        _ph("No"),
        _ph("Nama Barang", TA_LEFT),
        _ph("Kategori", TA_LEFT),
        _ph("Jml & Satuan"),
        _ph("Harga Satuan", TA_RIGHT),
        _ph("Total", TA_RIGHT),
    ]]

    # Baris data item
    baris_data: list[list[Any]] = []
    for idx, item in enumerate(daftar_item):
        jml_sat = f"{format_jumlah(item['jumlah_barang'])} {item['nama_satuan']}"
        baris_data.append([
            _pd(str(idx + 1), TA_CENTER),
            _pd(item["nama_barang"]),
            _pd(item["nama_kategori"]),
            _pd(jml_sat, TA_CENTER),
            _pd(format_rupiah(item["harga_satuan"]), TA_RIGHT),
            _pd(format_rupiah(item["harga_total"]), TA_RIGHT),
        ])

    # Baris footer total
    total_semua = sum(float(item["harga_total"]) for item in daftar_item)
    baris_footer: list[list[Any]] = [[
        _pd(""), _pd(""), _pd(""), _pd(""),
        _pd("Total:", TA_RIGHT, bold=True, warna=PDF_TEAL_900),
        _pd(format_rupiah(total_semua), TA_RIGHT, bold=True, warna=PDF_TEAL_900),
    ]]

    # Jika tidak ada item, tampilkan pesan kosong
    if not daftar_item:
        baris_data = [[
            Paragraph(
                "Tidak ada item belanja pada periode ini.",
                _gaya("empty", "Helvetica", 9, PDF_TEKS_ABU, TA_CENTER),
            ),
            _pd(""), _pd(""), _pd(""), _pd(""), _pd(""),
        ]]

    semua_baris = baris_header + baris_data + baris_footer
    n_isi = len(baris_data)

    tbl_item = Table(semua_baris, colWidths=lebar_kolom, repeatRows=1)

    style_item: list[Any] = [
        # Header (baris 0)
        ("BACKGROUND",    (0, 0), (-1,  0), PDF_TEAL_600),
        ("TOPPADDING",    (0, 0), (-1,  0), 7),
        ("BOTTOMPADDING", (0, 0), (-1,  0), 7),
        ("LEFTPADDING",   (0, 0), (-1,  0), 5),
        ("RIGHTPADDING",  (0, 0), (-1,  0), 5),
        # Baris data (baris 1 s.d. sebelum footer)
        ("TOPPADDING",    (0, 1), (-1, -2), 6),
        ("BOTTOMPADDING", (0, 1), (-1, -2), 6),
        ("LEFTPADDING",   (0, 1), (-1, -2), 5),
        ("RIGHTPADDING",  (0, 1), (-1, -2), 5),
        # Footer (baris terakhir)
        ("BACKGROUND",    (0, -1), (-1, -1), PDF_TEAL_100),
        ("TOPPADDING",    (0, -1), (-1, -1), 7),
        ("BOTTOMPADDING", (0, -1), (-1, -1), 7),
        ("LEFTPADDING",   (0, -1), (-1, -1), 5),
        ("RIGHTPADDING",  (0, -1), (-1, -1), 5),
        ("LINEABOVE",     (0, -1), (-1, -1), 1, PDF_TEAL_600),
        # Border
        ("LINEBELOW",     (0,  0), (-1,  0), 0.5, PDF_TEAL_200),
        ("GRID",          (0,  1), (-1, -2), 0.3, PDF_TEAL_200),
        ("BOX",           (0,  0), (-1, -1), 0.5, PDF_TEAL_200),
        # Vertical align tengah semua sel
        ("VALIGN",        (0,  0), (-1, -1), "MIDDLE"),
    ]

    # Warnai baris genap dengan teal-50 (alternating rows)
    for i in range(2, n_isi + 1):
        if i % 2 == 0:
            style_item.append(("BACKGROUND", (0, i), (-1, i), PDF_TEAL_50))

    tbl_item.setStyle(TableStyle(style_item))
    elemen.append(tbl_item)

    # ── 5. TIMESTAMP ─────────────────────────────────────────────────────────
    elemen.append(Spacer(1, 18))
    elemen.append(HRFlowable(
        width="100%", thickness=0.5,
        color=PDF_TEAL_200, spaceAfter=6,
    ))
    elemen.append(Paragraph(
        f"Diekspor pada: {timestamp}",
        _gaya("ts", "Helvetica", 7.5, PDF_TEKS_ABU, TA_LEFT),
    ))

    doc.build(elemen, canvasmaker=_NumberedCanvas)


# ─── EKSPOR EXCEL ─────────────────────────────────────────────────────────────

def ekspor_excel(path_file: str, data_periode: dict[str, Any], daftar_item: list[dict[str, Any]]) -> None:
    """
    Generate file Excel (.xlsx) laporan belanja dan simpan ke path_file.
    Parameter identik dengan ekspor_pdf().
    """

    def _fmt(tanggal: str) -> str:
        return "-".join(reversed(tanggal.split("-")))

    nama_bulan   = NAMA_BULAN[data_periode["bulan"] - 1]
    tgl_mulai    = _fmt(data_periode["tgl_mulai"])
    tgl_selesai  = _fmt(data_periode["tgl_selesai"])
    sisa         = float(data_periode["sisa_budget"])
    nominal      = float(data_periode["nominal_budget"])
    timestamp    = datetime.now().strftime("%d-%m-%Y %H:%M")

    wb = Workbook()

    # wb.active mengembalikan Optional[Worksheet]; pastikan tidak None
    ws_raw = wb.active
    if ws_raw is None:
        # Sangat tidak mungkin terjadi tapi perlu untuk tipe yang benar
        ws_raw = wb.create_sheet()
    ws: Worksheet = ws_raw
    ws.title = f"Laporan Minggu ke-{data_periode['minggu_ke']}"

    # ── Helper style lokal ──────────────────────────────────────────────────

    def _fill(hex_warna: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_warna)

    def _font(bold: bool = False, ukuran: int = 11,
              warna: str = "000000", italic: bool = False) -> Font:
        return Font(bold=bold, size=ukuran, color=warna, italic=italic, name="Calibri")

    def _align(horizontal: str = "left", vertical: str = "center",
               wrap: bool = False, indent: int = 0) -> Alignment:
        return Alignment(horizontal=horizontal, vertical=vertical,
                         wrap_text=wrap, indent=indent)

    def _border(warna: str = XL_TEAL_200, tebal: str = "thin") -> Border:
        sisi = Side(style=tebal, color=warna)  # type: ignore[arg-type]
        return Border(left=sisi, right=sisi, top=sisi, bottom=sisi)

    def _set(baris: int, kolom: int, nilai: Any,
             bold: bool = False, ukuran: int = 11, warna_font: str = "000000",
             fill: str | None = None, align_h: str = "left",
             italic: bool = False, num_fmt: str | None = None,
             wrap: bool = False, indent: int = 0) -> Any:
        """Isi satu sel Excel dengan semua atribut style sekaligus."""
        sel = ws.cell(row=baris, column=kolom, value=nilai)
        sel.font      = _font(bold, ukuran, warna_font, italic)
        sel.alignment = _align(align_h, wrap=wrap, indent=indent)
        if fill is not None:
            sel.fill = _fill(fill)
        if num_fmt is not None:
            sel.number_format = num_fmt
        return sel

    # ── Lebar kolom ──────────────────────────────────────────────────────────
    # A=No, B=Nama, C=Kategori, D=Jumlah, E=Satuan, F=Harga Satuan, G=Total
    lebar_kol = {"A": 5, "B": 28, "C": 20, "D": 9, "E": 11, "F": 16, "G": 16}
    for huruf, lebar in lebar_kol.items():
        ws.column_dimensions[huruf].width = lebar

    # ── Row 1: Judul ──────────────────────────────────────────────────────────
    ws.merge_cells("A1:G1")
    _set(1, 1, "CATATAN BELANJA",
         bold=True, ukuran=16, warna_font=XL_PUTIH,
         fill=XL_TEAL_900, align_h="left", indent=1)
    ws.row_dimensions[1].height = 30

    # ── Row 2: Sub-judul periode ──────────────────────────────────────────────
    ws.merge_cells("A2:G2")
    _set(2, 1,
         f"Laporan Minggu ke-{data_periode['minggu_ke']}  \u00b7  {nama_bulan} {data_periode['tahun']}",
         ukuran=10, warna_font=XL_TEAL_200,
         fill=XL_TEAL_900, align_h="left", indent=1)
    ws.row_dimensions[2].height = 16

    # ── Row 3: Rentang tanggal ────────────────────────────────────────────────
    ws.merge_cells("A3:G3")
    _set(3, 1, f"{tgl_mulai}  \u2013  {tgl_selesai}",
         ukuran=9, warna_font=XL_TEAL_200,
         fill=XL_TEAL_900, align_h="left", indent=1)
    ws.row_dimensions[3].height = 14

    # ── Row 4: Pemisah kosong ────────────────────────────────────────────────
    ws.row_dimensions[4].height = 8

    # ── Row 5: Label stat (Budget | Total Pengeluaran | Sisa Budget) ─────────
    ws.merge_cells("A5:B5")
    _set(5, 1, "Budget",
         ukuran=8, warna_font=XL_TEAL_200, fill=XL_TEAL_900, align_h="center")
    ws.row_dimensions[5].height = 13

    ws.merge_cells("C5:D5")
    _set(5, 3, "Total Pengeluaran",
         ukuran=8, warna_font=XL_TEAL_200, fill=XL_TEAL_900, align_h="center")

    ws.merge_cells("E5:G5")
    _set(5, 5, "Sisa Budget",
         ukuran=8, warna_font=XL_TEAL_200, fill=XL_TEAL_900, align_h="center")

    # ── Row 6: Nilai stat ─────────────────────────────────────────────────────
    ws.merge_cells("A6:B6")
    _set(6, 1, format_rupiah(nominal),
         bold=True, ukuran=12, warna_font=XL_PUTIH,
         fill=XL_TEAL_900, align_h="center")
    ws.row_dimensions[6].height = 22

    ws.merge_cells("C6:D6")
    _set(6, 3, format_rupiah(data_periode["total_pengeluaran"]),
         bold=True, ukuran=12, warna_font=XL_PUTIH,
         fill=XL_TEAL_900, align_h="center")

    ws.merge_cells("E6:G6")
    # Tentukan warna teks sisa budget sesuai threshold
    if sisa < 0:
        warna_sisa_xl = XL_DANGER
    elif nominal > 0 and sisa < nominal * 0.20:
        warna_sisa_xl = XL_WARNING
    else:
        warna_sisa_xl = XL_PUTIH
    _set(6, 5, format_rupiah(sisa),
         bold=True, ukuran=12, warna_font=warna_sisa_xl,
         fill=XL_TEAL_900, align_h="center")

    # ── Row 7: Pemisah sebelum tabel ─────────────────────────────────────────
    ws.row_dimensions[7].height = 10

    # ── Row 8: Header kolom tabel ─────────────────────────────────────────────
    header_kolom = [
        ("No", "center"), ("Nama Barang", "left"), ("Kategori", "left"),
        ("Jumlah", "center"), ("Satuan", "center"),
        ("Harga Satuan", "right"), ("Harga Total", "right"),
    ]
    for kolom, (teks, align_h) in enumerate(header_kolom, start=1):
        sel_h = _set(8, kolom, teks,
                     bold=True, ukuran=9, warna_font=XL_PUTIH,
                     fill=XL_TEAL_600, align_h=align_h)
        sel_h.border = _border(XL_TEAL_600, "medium")
    ws.row_dimensions[8].height = 18

    # ── Row 9+: Data item ─────────────────────────────────────────────────────
    for idx, item in enumerate(daftar_item):
        baris    = 9 + idx
        # Baris genap (idx 1, 3, 5, ...) → warna alternating teal-50
        fill_alt: str | None = XL_TEAL_50 if idx % 2 == 1 else None

        kolom_data = [
            (str(idx + 1),               "center", None),
            (item["nama_barang"],          "left",  None),
            (item["nama_kategori"],         "left",  None),
            (item["jumlah_barang"],         "center", "0.##"),
            (item["nama_satuan"],           "center", None),
            (item["harga_satuan"],          "right",  "#,##0"),
            (item["harga_total"],           "right",  "#,##0"),
        ]
        for kolom, (nilai, align_h, num_fmt) in enumerate(kolom_data, start=1):
            sel_d = _set(baris, kolom, nilai,
                         ukuran=9, fill=fill_alt, align_h=align_h,
                         num_fmt=num_fmt)
            sel_d.border = _border()
        ws.row_dimensions[baris].height = 16

    # ── Footer total ──────────────────────────────────────────────────────────
    baris_footer = 9 + len(daftar_item)
    total_semua  = sum(float(item["harga_total"]) for item in daftar_item)

    # Warnai semua sel di baris footer
    for kolom in range(1, 8):
        sel_f = ws.cell(row=baris_footer, column=kolom)
        sel_f.fill   = _fill(XL_TEAL_100)
        sel_f.border = _border(XL_TEAL_600)

    # Tulis label "Total:" di kolom F
    sel_lbl = _set(baris_footer, 6, "Total:",
                   bold=True, ukuran=9, warna_font=XL_TEKS,
                   fill=XL_TEAL_100, align_h="right")
    sel_lbl.border = _border(XL_TEAL_600)

    # Tulis nilai total di kolom G dengan format angka native Excel
    sel_tot = _set(baris_footer, 7, total_semua,
                   bold=True, ukuran=9, warna_font=XL_TEKS,
                   fill=XL_TEAL_100, align_h="right", num_fmt="#,##0")
    sel_tot.border = _border(XL_TEAL_600)
    ws.row_dimensions[baris_footer].height = 18

    # ── Timestamp ekspor ──────────────────────────────────────────────────────
    baris_ts = baris_footer + 2
    ws.merge_cells(f"A{baris_ts}:G{baris_ts}")
    sel_ts = ws.cell(row=baris_ts, column=1,
                     value=f"Diekspor pada: {timestamp}")
    sel_ts.font      = Font(italic=True, size=8, color=XL_ABU, name="Calibri")
    sel_ts.alignment = Alignment(horizontal="left", vertical="center")

    wb.save(path_file)